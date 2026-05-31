import openpyxl
import sys

try:
    sys.stdout.reconfigure(encoding='utf-8')
except:
    pass

# Load Excel file
wb = openpyxl.load_workbook(r'C:\Code\Sup_AI_Edu\assets\Giảng viên\danh_sach_giang_vien_day_du_704_dong.xlsx')
ws = wb.active

print(f'📊 Sheet name: {ws.title}')
print(f'📈 Total rows: {ws.max_row}')
print(f'📋 Total columns: {ws.max_column}')
print()
print('🔍 Column headers:')
for col in range(1, ws.max_column + 1):
    header = ws.cell(1, col).value
    print(f'  Col {col}: {header}')

print()
print('📝 First 3 rows of data:')
for row in range(2, min(5, ws.max_row + 1)):
    row_data = []
    for col in range(1, ws.max_column + 1):
        cell_value = ws.cell(row, col).value
        row_data.append(str(cell_value)[:25] if cell_value else 'N/A')
    print(f'  Row {row}: {" | ".join(row_data)}')
